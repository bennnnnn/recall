"""Attachment MIME allowlists and text extraction for chat context."""

from __future__ import annotations

import asyncio
import io
import logging
from dataclasses import dataclass
from typing import Any
from uuid import UUID

from sqlalchemy.ext.asyncio import AsyncSession

from app.core.attachment_limits import MAX_ATTACHMENT_SIZE as MAX_ATTACHMENT_SIZE
from app.core.config import Settings
from app.gateways.storage_gateway import StorageGateway
from app.modules.attachments.content_types import (
    _DOCX_CONTENT_TYPE,
    _OFFICE_MARKERS,
    _TEXTISH_TYPES,
    EXTRACTABLE_CONTENT_TYPES,
    _docx_zip_bomb,
    _office_content_type,
    bytes_match_claimed,
)

# Re-exported: rag and chat's vision paths read it through this module.
from app.modules.attachments.content_types import is_image_content_type as is_image_content_type
from app.modules.attachments.office_text import extract_office_text

logger = logging.getLogger(__name__)

# Inline excerpt cap — keeps the user-message attachment excerpt small so it
# doesn't bloat the prompt with a full document dump.
MAX_EXTRACT_CHARS = 12_000
# Indexing cap — the background indexing path chunks the whole document for
# pgvector RAG, so it can extract far more than the inline excerpt. Capped to
# avoid runaway memory on adversarially large text payloads.
MAX_INDEX_EXTRACT_CHARS = 200_000
# Inline excerpts stay small; background indexing can read later PDF pages.
PDF_EXTRACT_MAX_PAGES = 25
PDF_INDEX_MAX_PAGES = 500


class ExtractTimeoutError(TimeoutError):
    """Parsing ran past its budget; the file may still have text."""


@dataclass(frozen=True)
class ExtractedText:
    text: str
    char_capped: bool = False
    page_capped: bool = False
    via_ocr: bool = False


def file_excerpt_limit_note(
    *,
    char_capped: bool,
    page_capped: bool,
    max_chars: int = MAX_EXTRACT_CHARS,
    max_pages: int = PDF_EXTRACT_MAX_PAGES,
) -> str | None:
    """Model-facing note when only a prefix of the file was read. Hidden in the UI
    because it sits inside the ``[File (type)]`` excerpt block."""
    if not char_capped and not page_capped:
        return None
    bits: list[str] = []
    if page_capped:
        bits.append(f"the first {max_pages} pages")
    if char_capped:
        bits.append(f"the first {max_chars} characters")
    covered = " and ".join(bits)
    return (
        f"[File note: only {covered} were read; later pages or text may be missing. "
        "Do not claim facts from unread parts. If asked about later content, say "
        "it was not in the excerpt.]"
    )


def _docx_tables_as_markdown(document: Any) -> list[str]:
    """Keep table cell alignment as a markdown grid instead of flattening cells."""
    blocks: list[str] = []
    for table in document.tables:
        rows_out: list[str] = []
        for row in table.rows:
            cells = [cell.text.strip().replace("|", "\\|") for cell in row.cells]
            if not any(cells):
                continue
            rows_out.append("| " + " | ".join(cells) + " |")
        if not rows_out:
            continue
        col_count = max(1, rows_out[0].count("|") - 1)
        sep = "| " + " | ".join("---" for _ in range(col_count)) + " |"
        if len(rows_out) == 1:
            blocks.append(f"{rows_out[0]}\n{sep}")
        else:
            blocks.append(f"{rows_out[0]}\n{sep}\n" + "\n".join(rows_out[1:]))
    return blocks


def extract_text_details(
    content_type: str,
    data: bytes,
    *,
    max_chars: int = MAX_EXTRACT_CHARS,
) -> ExtractedText | None:
    """Sync, CPU-bound parsing with truncation flags for honest disclosure.

    Call via ``extract_text_details_async`` on any code path that isn't already
    off the event loop.
    """
    if content_type in _TEXTISH_TYPES:
        text = data.decode("utf-8", errors="replace").strip()
        if not text:
            return None
        return ExtractedText(text=text[:max_chars], char_capped=len(text) > max_chars)

    if content_type == "application/pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            page_limit = (
                PDF_INDEX_MAX_PAGES if max_chars > MAX_EXTRACT_CHARS else PDF_EXTRACT_MAX_PAGES
            )
            page_capped = len(reader.pages) > page_limit
            parts: list[str] = []
            size = 0
            char_capped = False
            for index, page in enumerate(reader.pages[:page_limit], start=1):
                page_text = page.extract_text().replace("\x00", "")  # Postgres text rejects NUL
                if page_text:
                    parts.append(f"[page {index}] {page_text.strip()}")
                    size += len(parts[-1]) + 2
                    if size > max_chars:
                        char_capped = True
                        break
            joined = "\n\n".join(parts).strip()
            if not joined:
                return None
            return ExtractedText(
                text=joined[:max_chars],
                char_capped=char_capped or len(joined) > max_chars,
                page_capped=page_capped,
            )
        except Exception:
            logger.warning(
                "PDF text extraction failed; the file has no text for search", exc_info=True
            )
            return None

    if content_type == _DOCX_CONTENT_TYPE:
        try:
            if _docx_zip_bomb(data):
                return None
            from docx import Document

            document = Document(io.BytesIO(data))
            parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
            parts.extend(_docx_tables_as_markdown(document))
            joined = "\n".join(parts).strip()
            if not joined:
                return None
            return ExtractedText(text=joined[:max_chars], char_capped=len(joined) > max_chars)
        except Exception:
            logger.warning(
                "DOCX text extraction failed; the file has no text for search", exc_info=True
            )
            return None

    if content_type in _OFFICE_MARKERS.values() and content_type != _DOCX_CONTENT_TYPE:
        if _office_content_type(data) != content_type:
            return None
        office_text = extract_office_text(content_type, data, max_chars=max_chars + 1)
        if office_text:
            return ExtractedText(
                text=office_text[:max_chars], char_capped=len(office_text) > max_chars
            )

    # Legacy .doc (application/msword) and anything else: no parser. Handled
    # explicitly (not silently) by format_attachment_lines via
    # EXTRACTABLE_CONTENT_TYPES.
    return None


async def extract_text_details_async(
    content_type: str,
    data: bytes,
    settings: Settings,
    *,
    max_chars: int = MAX_EXTRACT_CHARS,
    allow_ocr: bool = True,
    ocr_max_pages: int | None = None,
    timeout_seconds: float | None = None,
    raise_on_timeout: bool = False,
) -> ExtractedText | None:
    """Offload the sync, CPU-bound parse to a worker thread with a timeout,
    so a large or adversarially crafted PDF/DOCX can't block the event loop —
    same pattern as the SymPy math solve offload.

    When a PDF has no text layer, fall through to vision OCR (scanned /
    photographed pages) using a separate network timeout — unless
    *allow_ocr* is False (turn prepare; OCR belongs on the index job).
    A CPU-extract timeout does not then spend another 30s on vision.
    ``timeout_seconds`` defaults to the turn's short budget; the index job
    passes its own and ``raise_on_timeout`` (a slow parse is not "no text").
    """
    timed_out = False
    try:
        async with asyncio.timeout(timeout_seconds or settings.attachment_extract_timeout_seconds):
            details = await asyncio.to_thread(
                extract_text_details, content_type, data, max_chars=max_chars
            )
    except TimeoutError as exc:
        logger.warning("Attachment text extraction timed out for content_type=%s", content_type)
        if raise_on_timeout:
            raise ExtractTimeoutError(content_type) from exc
        timed_out = True
        details = None
    if details:
        return details
    if timed_out:
        return None
    if allow_ocr and content_type == "application/pdf" and settings.attachment_ocr_enabled:
        from app.modules.attachments.ocr import ocr_scanned_pdf

        pages_for_ocr = (
            ocr_max_pages
            if ocr_max_pages is not None
            else (
                settings.attachment_ocr_index_max_pages
                if max_chars > MAX_EXTRACT_CHARS
                else settings.attachment_ocr_max_pages
            )
        )
        ocr_result = await ocr_scanned_pdf(
            settings,
            data,
            max_chars=max_chars,
            max_pages=pages_for_ocr,
        )
        if ocr_result is None:
            return None
        if isinstance(ocr_result, ExtractedText):
            return ocr_result
        return ExtractedText(text=ocr_result, via_ocr=True)
    return None


async def read_attachment_bytes(gateway: StorageGateway, storage_key: str) -> bytes | None:
    return await gateway.read_bytes(storage_key)


async def verify_uploaded_bytes(
    gateway: StorageGateway,
    *,
    content_type: str,
    storage_key: str,
    max_size: int = MAX_ATTACHMENT_SIZE,
    declared_size: int | None = None,
) -> tuple[bytes | None, str | None]:
    """Return uploaded bytes when they match the declared type, else an error detail.

    When *declared_size* is provided, the actual byte count must match exactly
    — a mismatch means the client lied about the size at presign time (or the
    upload was truncated/extended), and the stored object is not what the DB
    row claims. Without this check, a presign for a 1-byte "image/png" could
    be used to store a 10 MB blob that later gets served back as a PNG.
    """
    data = await read_attachment_bytes(gateway, storage_key)
    if not data:
        return None, "Upload not found"
    if len(data) > max_size:
        return None, "Invalid upload size"
    if declared_size is not None and len(data) != declared_size:
        return None, "Uploaded size does not match the declared size"
    if not bytes_match_claimed(content_type, data):
        return None, "Uploaded bytes do not match the declared content type"
    return data, None


async def purge_invalid_upload(
    gateway: StorageGateway,
    session: AsyncSession,
    *,
    attachment_id: UUID,
    storage_key: str,
) -> bool:
    from app.modules.attachments import repository as attachments_repo

    removed = await attachments_repo.delete_rows(session, [attachment_id], commit=False)
    # Commit row deletion before irreversible object removal. Failed storage
    # cleanup must survive the loss of its attachment row.
    await session.commit()
    if not removed:
        return False
    try:
        await gateway.delete_bytes(storage_key)
    except Exception:
        from app.modules.attachments.lifecycle import enqueue_failed_storage_deletes

        await enqueue_failed_storage_deletes([storage_key])
    return True


async def ensure_verified_or_purge(
    gateway: StorageGateway,
    session: AsyncSession,
    *,
    attachment_id: UUID,
    content_type: str,
    storage_key: str,
    declared_size: int | None = None,
) -> tuple[str | None, bool]:
    """Verify R2/S3 bytes match declared type; purge row+object on failure.

    Local dev uploads are validated on PUT /upload — no-op for that backend.
    Returns the error detail and whether this call actually purged the row.
    """
    from app.gateways.storage_gateway import LocalStorageGateway
    from app.modules.attachments import repository as attachments_repo

    if isinstance(gateway, LocalStorageGateway):
        return None, False
    _, error = await verify_uploaded_bytes(
        gateway,
        content_type=content_type,
        storage_key=storage_key,
        declared_size=declared_size,
    )
    if error:
        removed = await purge_invalid_upload(
            gateway,
            session,
            attachment_id=attachment_id,
            storage_key=storage_key,
        )
        return error, removed
    await attachments_repo.mark_verified(session, attachment_id)
    return None, False


def _scanned_pdf_note(settings: Settings) -> str:
    """What the model may promise about a PDF with no text layer: only what will run."""
    lead = (
        "[File attached: application/pdf. No extractable text layer — this is likely a "
        "scanned or image-only PDF. "
    )
    if settings.attachment_ocr_enabled and settings.attachment_rag_enabled:
        return (
            f"{lead}Recall will try to read its pages in the background, within the user's "
            "daily photo allowance; if that works, the text can be searched on later turns. "
            "Do not promise it; suggest a text-based PDF or pasting the text if needed.]"
        )
    return f"{lead}Recall cannot read scanned pages; suggest a text-based PDF or pasting the text.]"


async def format_attachment_lines(
    gateway: StorageGateway,
    *,
    attachment_id: str,
    content_type: str,
    storage_key: str,
    size_bytes: int,
    settings: Settings,
    data: bytes | None = None,
) -> tuple[list[str], bool]:
    """Return prompt lines and whether the attachment is an image.

    When *data* is provided (e.g. bytes already loaded by turn-prep verify),
    skip a second storage download for text extraction.
    """
    if is_image_content_type(content_type):
        # API path the mobile app can fetch with auth (stored in message content).
        return [f"[Image: /attachments/{attachment_id}/file]"], True

    if data is None:
        data = await read_attachment_bytes(gateway, storage_key)
    file_ref = f"[File: /attachments/{attachment_id}/file]"
    if data:
        excerpt = await extract_text_details_async(content_type, data, settings, allow_ocr=False)
        if excerpt:
            note = file_excerpt_limit_note(
                char_capped=excerpt.char_capped,
                page_capped=excerpt.page_capped,
            )
            body = f"{note}\n{excerpt.text}" if note else excerpt.text
            return [file_ref, f"[File ({content_type})]\n{body}"], False
        if content_type not in EXTRACTABLE_CONTENT_TYPES:
            return (
                [
                    file_ref,
                    f"[File attached: {content_type}. Recall can't read this file type yet — "
                    "tell the user to try a PDF, Word (.docx), or plain text file instead.]",
                ],
                False,
            )
        if content_type == "application/pdf":
            return [file_ref, _scanned_pdf_note(settings)], False
        return (
            [
                file_ref,
                f"[File attached: {content_type}. No extractable text was found — "
                "tell the user the file appears empty or unreadable.]",
            ],
            False,
        )

    return [file_ref, f"[File attached: {content_type}, {size_bytes} bytes]"], False
