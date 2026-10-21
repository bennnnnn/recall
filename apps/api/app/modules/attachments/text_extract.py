"""A document's text, parsed: what a document worker runs (core/document_executor.py).

Sync and CPU-bound (pypdf, python-docx, the office readers). Split from ``content`` so a
worker imports only the parsers, not the database or storage stack.
"""

from __future__ import annotations

import io
import logging
from dataclasses import dataclass
from typing import Any

from app.modules.attachments.content_types import (
    _DOCX_CONTENT_TYPE,
    _OFFICE_MARKERS,
    _TEXTISH_TYPES,
    _docx_zip_bomb,
    _office_content_type,
)
from app.modules.attachments.office_text import extract_office_text

logger = logging.getLogger(__name__)

# Inline excerpt cap — keeps the user-message attachment excerpt small so it
# doesn't bloat the prompt with a full document dump.
MAX_EXTRACT_CHARS = 12_000


# Indexing cap — the background indexing path chunks the whole document for
# pgvector RAG, so it can extract far more than the inline excerpt. Capped to
# avoid runaway memory on adversarially large text payloads.
MAX_INDEX_EXTRACT_CHARS = 200_000


# Inline excerpts stay small; background indexing can read later PDF pages.
PDF_EXTRACT_MAX_PAGES = 25


PDF_INDEX_MAX_PAGES = 500


@dataclass(frozen=True)
class ExtractedText:
    text: str
    char_capped: bool = False
    page_capped: bool = False
    via_ocr: bool = False


def _docx_tables_as_markdown(document: Any) -> list[str]:
    """Keep table cell alignment as a markdown grid instead of flattening cells."""
    blocks: list[str] = []
    for table in document.tables:
        rows_out: list[str] = []
        for row in table.rows:
            cells = [cell.text.strip().replace("|", "\\|") for cell in row.cells]
            if not any(cells):
                continue
            rows_out.append("| " + " | ".join(cells) + " |")
        if not rows_out:
            continue
        col_count = max(1, rows_out[0].count("|") - 1)
        sep = "| " + " | ".join("---" for _ in range(col_count)) + " |"
        if len(rows_out) == 1:
            blocks.append(f"{rows_out[0]}\n{sep}")
        else:
            blocks.append(f"{rows_out[0]}\n{sep}\n" + "\n".join(rows_out[1:]))
    return blocks


def extract_text_details(
    content_type: str,
    data: bytes,
    *,
    max_chars: int = MAX_EXTRACT_CHARS,
) -> ExtractedText | None:
    """Sync, CPU-bound parsing with truncation flags for honest disclosure.

    Call via ``extract_text_details_async`` on any code path that isn't already
    off the event loop.
    """
    if content_type in _TEXTISH_TYPES:
        text = data.decode("utf-8", errors="replace").strip()
        if not text:
            return None
        return ExtractedText(text=text[:max_chars], char_capped=len(text) > max_chars)

    if content_type == "application/pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            page_limit = (
                PDF_INDEX_MAX_PAGES if max_chars > MAX_EXTRACT_CHARS else PDF_EXTRACT_MAX_PAGES
            )
            page_capped = len(reader.pages) > page_limit
            parts: list[str] = []
            size = 0
            char_capped = False
            for index, page in enumerate(reader.pages[:page_limit], start=1):
                page_text = page.extract_text().replace("\x00", "")  # Postgres text rejects NUL
                if page_text:
                    parts.append(f"[page {index}] {page_text.strip()}")
                    size += len(parts[-1]) + 2
                    if size > max_chars:
                        char_capped = True
                        break
            joined = "\n\n".join(parts).strip()
            if not joined:
                return None
            return ExtractedText(
                text=joined[:max_chars],
                char_capped=char_capped or len(joined) > max_chars,
                page_capped=page_capped,
            )
        except Exception:
            logger.warning(
                "PDF text extraction failed; the file has no text for search", exc_info=True
            )
            return None

    if content_type == _DOCX_CONTENT_TYPE:
        try:
            if _docx_zip_bomb(data):
                return None
            from docx import Document

            document = Document(io.BytesIO(data))
            parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
            parts.extend(_docx_tables_as_markdown(document))
            joined = "\n".join(parts).strip()
            if not joined:
                return None
            return ExtractedText(text=joined[:max_chars], char_capped=len(joined) > max_chars)
        except Exception:
            logger.warning(
                "DOCX text extraction failed; the file has no text for search", exc_info=True
            )
            return None

    if content_type in _OFFICE_MARKERS.values() and content_type != _DOCX_CONTENT_TYPE:
        if _office_content_type(data) != content_type:
            return None
        office_text = extract_office_text(content_type, data, max_chars=max_chars + 1)
        if office_text:
            return ExtractedText(
                text=office_text[:max_chars], char_capped=len(office_text) > max_chars
            )

    # Legacy .doc (application/msword) and anything else: no parser. Handled
    # explicitly (not silently) by format_attachment_lines via
    # EXTRACTABLE_CONTENT_TYPES.
    return None


def extract_in_worker(content_type: str, data: bytes, max_chars: int) -> ExtractedText | None:
    """``extract_text_details`` with positional arguments, as a pool job takes them."""
    return extract_text_details(content_type, data, max_chars=max_chars)
